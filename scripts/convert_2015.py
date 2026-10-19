import json
import time
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from _common import find_file, get_paths

FILES = {
    "normal_v0": "SWaT_Dataset_Normal_v0.xlsx",
    "normal_v1": "SWaT_Dataset_Normal_v1.xlsx",
    "attack_v0": "SWaT_Dataset_Attack_v0.xlsx",
}
TS_FORMATS = ["%d/%m/%Y %I:%M:%S %p",
              "%d/%m/%Y %H:%M:%S", "%m/%d/%Y %I:%M:%S %p"]


def parse_ts(v):
    if isinstance(v, datetime):
        return v
    s = str(v).strip()
    for f in TS_FORMATS:
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def to_float(v):
    if v is None:
        return np.nan
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def read_sheet(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    max_row = ws.max_row or 600_000
    header, cols, ts, lab, data, bad_vals, i = None, None, [], [], None, 0, 0
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] is not None and str(row[0]).strip().lower() == "timestamp":
                header = [None if c is None else str(c).strip() for c in row]
                cols = header[1:-1]
                data = np.full((max_row, len(cols)), np.nan, dtype=np.float64)
            continue
        if row is None or all(v is None for v in row):
            continue
        ts.append(parse_ts(row[0]))
        lab.append(None if row[-1] is None else str(row[-1]))
        vals = [to_float(v) for v in row[1:-1]]
        bad_vals += sum(1 for v, r in zip(vals,
                        row[1:-1]) if np.isnan(v) and r is not None)
        data[i, :] = vals
        i += 1
    wb.close()
    df = pd.DataFrame(data[:i], columns=cols).astype("float32")
    df.insert(0, "timestamp", pd.to_datetime(pd.Series(ts)))
    raw_labels = pd.Series(lab)
    df["label_raw"] = raw_labels
    df["attack"] = raw_labels.str.replace(
        " ", "", regex=False).str.strip().str.lower().eq("attack").astype("int8")
    return df, header, bad_vals


def segments(flag, ts):
    f = flag.to_numpy()
    edges = np.flatnonzero(np.diff(np.r_[0, f, 0]))
    out = []
    for s, e in zip(edges[::2], edges[1::2]):
        out.append({"start": str(ts.iloc[s]), "end": str(
            ts.iloc[e - 1]), "seconds": int(e - s)})
    return out


def main():
    root, work = get_paths(__doc__)
    (work / "parquet").mkdir(exist_ok=True)
    (work / "reports").mkdir(exist_ok=True)
    report, stats = {}, []
    for key, name in FILES.items():
        t0 = time.time()
        path = find_file(root, name)
        print(f"Reading {name} ...")
        df, header, bad_vals = read_sheet(path)
        d = df["timestamp"].diff().dt.total_seconds()
        sensors = [c for c in df.columns if c not in (
            "timestamp", "label_raw", "attack")]
        r = {
            "file": str(path.relative_to(root)),
            "rows": int(len(df)),
            "columns_raw_header": header,
            "n_sensor_columns": len(sensors),
            "first_ts": str(df["timestamp"].iloc[0]),
            "last_ts": str(df["timestamp"].iloc[-1]),
            "unparsed_timestamps": int(df["timestamp"].isna().sum()),
            "non_numeric_values": int(bad_vals),
            "missing_values": int(df[sensors].isna().sum().sum()),
            "duplicate_timestamps": int(df["timestamp"].duplicated().sum()),
            "gaps_over_1s": int((d > 1).sum()),
            "steps_not_1s": {str(k): int(v) for k, v in d.value_counts().head(10).items()},
            "label_values_raw": {str(k): int(v) for k, v in df["label_raw"].value_counts(dropna=False).items()},
            "attack_seconds": int(df["attack"].sum()),
            "constant_columns": [c for c in sensors if df[c].nunique(dropna=True) <= 1],
            "seconds_to_read": round(time.time() - t0),
        }
        if df["attack"].any():
            segs = segments(df["attack"], df["timestamp"])
            r["attack_segments"] = segs
            r["n_attack_segments"] = len(segs)
        report[key] = r
        desc = df[sensors].describe().T
        desc["n_unique"] = df[sensors].nunique()
        desc.insert(0, "file", key)
        stats.append(desc.reset_index().rename(columns={"index": "sensor"}))
        out = work / "parquet" / f"swat2015_{key}.parquet"
        df.to_parquet(out, index=False)
        print(f"  {len(df):,} rows -> {out.name}  ({r['seconds_to_read']}s)")

    with open(work / "reports" / "report_2015.json",
              "w", encoding="utf-8") as f:
        json.dump(report, f, indent=1)
    pd.concat(stats).to_csv(work / "reports" / "stats_2015.csv",
                            index=False, float_format="%.6g")
    with open(work / "reports" / "report_2015.txt",
              "w", encoding="utf-8") as f:
        for key, r in report.items():
            f.write(f"== {key}: {r['file']}\n")
            for k in ["rows", "n_sensor_columns", "first_ts",
                      "last_ts", "unparsed_timestamps",
                      "non_numeric_values", "missing_values",
                      "duplicate_timestamps", "gaps_over_1s",
                      "label_values_raw", "attack_seconds",
                      "n_attack_segments", "constant_columns"
                      ]:
                if k in r:
                    f.write(f"  {k}: {r[k]}\n")
            f.write("\n")
    print("\nDone.")


if __name__ == "__main__":
    main()
