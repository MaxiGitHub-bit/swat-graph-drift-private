import io
import json
import re
import zipfile

from openpyxl import load_workbook

from _common import get_paths

# do not open zips over 5 GB here (the 2017 file)
SKIP_IF_LARGER = 5 * 1024**3
MAX_MEMBER = 300 * 1024**2  # read members up to 300 MB into memory
PEEK = 4
MAGIC = {b"PK\x03\x04": "zip", b"PK\x05\x06": "empty zip",
         b"Rar!": "rar", b"7z\xbc\xaf": "7z",
         b"\x1f\x8b": "gzip", b"<!DO": "html (probably a failed download page)",
         b"<htm": "html (failed download page)",
         b"\xd0\xcf\x11\xe0": "old Office file"
         }


def magic_of(path):
    with open(path, "rb") as f:
        head = f.read(8)
    for sig, name in MAGIC.items():
        if head.startswith(sig):
            return name, head.hex()
    return "unknown", head.hex()


def cells(row):
    return [None if v is None else str(v) for v in row]


def peek_xlsx(data):
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows, n, last = [], 0, None
        for row in ws.iter_rows(values_only=True):
            if n < PEEK:
                rows.append(cells(row))
            last = row
            n += 1
        out.append({"sheet": ws.title, "dimension": ws.calculate_dimension(),
                    "n_rows": n, "first_rows": rows, "last_row": cells(last) if last else None}
                   )
    wb.close()
    return {"sheets": out}


def peek_text(data, max_lines=60):
    text = data.decode("utf-8", errors="replace")
    lines = text.splitlines()
    return {"n_lines": len(lines), "first_lines": lines[:max_lines], "last_line": lines[-1] if lines else ""}


def peek_docx(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    paras = [re.sub(r"<[^>]+>", "", p)
             for p in re.findall(r"<w:p[ >].*?</w:p>", xml)]
    return {"text": [p for p in paras if p.strip()][:80]}


def main():
    root, work = get_paths(__doc__)
    (work / "reports").mkdir(exist_ok=True)
    result = []
    for zp in sorted(root.rglob("*.zip")):
        rel = str(zp.relative_to(root))
        entry = {"zip": rel, "size_gb": round(zp.stat().st_size / 1024**3, 3)}
        kind, head = magic_of(zp)
        entry["file_type_by_content"] = kind
        if kind != "zip":
            entry["note"] = f"Not a normal zip (first bytes {head})."
            result.append(entry)
            continue
        if zp.stat().st_size > SKIP_IF_LARGER:
            entry["note"] = "Skipped here (too large); handled by a separate script."
            result.append(entry)
            continue
        members = []
        try:
            z = zipfile.ZipFile(zp)
        except zipfile.BadZipFile as e:
            entry["note"] = f"Zip header found but archive unreadable ({e}); download may be incomplete."
            result.append(entry)
            continue
        with z:
            for info in z.infolist():
                if info.is_dir():
                    continue
                m = {"name": info.filename, "size": info.file_size,
                     "compress_type": info.compress_type}
                low = info.filename.lower()
                try:
                    if low.rsplit("/", 1)[-1].startswith(("swat_dataset_normal", "swat_dataset_attack")):
                        m["note"] = "2015 file, handled by convert_2015.py"
                    elif info.file_size > MAX_MEMBER:
                        m["note"] = "too large to peek"
                    elif low.endswith((".xlsx", ".xlsm")):
                        m.update(peek_xlsx(z.read(info)))
                    elif low.endswith((".csv", ".txt", ".md")):
                        m.update(peek_text(z.read(info),
                                 60 if info.file_size < 200_000 else 6))
                    elif low.endswith(".docx"):
                        m.update(peek_docx(z.read(info)))
                except NotImplementedError as e:
                    m["error"] = f"compression not supported by Python ({e}); use 7-Zip to extract"
                except Exception as e:
                    m["error"] = f"{type(e).__name__}: {e}"
                members.append(m)
        entry["members"] = members
        result.append(entry)
        print(f"checked {rel}")

    with open(work / "reports" / "releases_peek.json", "w", encoding="utf-8") as f:
        json.dump(result, f, indent=1)
    with open(work / "reports" / "releases_peek.txt", "w", encoding="utf-8") as f:
        for e in result:
            f.write(
                f"\n### {e['zip']}  ({e['size_gb']} GB, {e['file_type_by_content']})\n")
            if "note" in e:
                f.write(f"NOTE: {e['note']}\n")
            for m in e.get("members", []):
                f.write(f"\n - {m['name']} [{m['size']:,} bytes]\n")
                for k in ("note", "error"):
                    if k in m:
                        f.write(f"{k.upper()}: {m[k]}\n")
                for s in m.get("sheets", []):
                    f.write(
                        f"sheet '{s['sheet']}' {s['dimension']}  rows={s['n_rows']:,}\n")
                    for r in s["first_rows"]:
                        f.write(f"{str(r)[:400]}\n")
                    f.write(f"LAST: {str(s['last_row'])[:400]}\n")
                if "first_lines" in m:
                    f.write(f"lines={m['n_lines']:,}\n")
                    for line in m["first_lines"][:12]:
                        f.write(f"{line[:400]}\n")
                    f.write(f"LAST: {m['last_line'][:400]}\n")
                for line in m.get("text", []):
                    f.write(f"{line[:400]}\n")
    print("\nDone, Victor.")


if __name__ == "__main__":
    main()
